"""Stable, ephemeral cloud API for the Smart Study Assistant portfolio demo."""

from __future__ import annotations

import json
import os
import re
import shutil
import threading
import time
import uuid
from collections import OrderedDict
from pathlib import Path
from typing import Any

import faiss
import gradio as gr
import numpy as np
import pdfplumber
import spaces
from huggingface_hub import InferenceClient
from pptx import Presentation
from sentence_transformers import SentenceTransformer

EMBEDDING_MODEL = "BAAI/bge-m3"
GENERATION_MODEL = "Qwen/Qwen3-4B-Instruct-2507"
CHUNK_SIZE, CHUNK_OVERLAP, TOP_K = 500, 75, 3
MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024
MAX_PDF_PAGES, MAX_PPTX_SLIDES = 60, 60
MAX_DOCUMENTS, DOCUMENT_TTL_SECONDS = 12, 45 * 60
MAX_GENERATION_CONTEXT_CHARS = 28_000
SESSION_ROOT = Path("/tmp/smart-study-assistant-sessions")

_documents: "OrderedDict[str, dict[str, Any]]" = OrderedDict()
_documents_lock = threading.RLock()
_embedding_models: dict[str, SentenceTransformer] = {}
_embedding_lock = threading.Lock()
_encoding_lock = threading.Lock()


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _extract_sections(file_path: str) -> tuple[list[dict[str, Any]], str]:
    if not file_path:
        raise ValueError("Upload a PDF or PPTX document first.")
    path = Path(file_path)
    if not path.is_file():
        raise ValueError("The uploaded file is no longer available. Upload it again.")
    if path.stat().st_size > MAX_FILE_SIZE_BYTES:
        raise ValueError("The file exceeds the 15 MB upload limit.")
    with path.open("rb") as stream:
        signature = stream.read(8)

    sections: list[dict[str, Any]] = []
    if signature.startswith(b"%PDF"):
        kind = "pdf"
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > MAX_PDF_PAGES:
                raise ValueError(f"PDFs are limited to {MAX_PDF_PAGES} pages.")
            for number, page in enumerate(pdf.pages, 1):
                text = (page.extract_text() or "").strip()
                if text:
                    sections.append({"text": text, "source_label": f"Page {number}", "page_number": number, "slide_number": None})
    elif signature.startswith(b"PK"):
        kind = "pptx"
        presentation = Presentation(str(path))
        if len(presentation.slides) > MAX_PPTX_SLIDES:
            raise ValueError(f"PowerPoint files are limited to {MAX_PPTX_SLIDES} slides.")
        for number, slide in enumerate(presentation.slides, 1):
            text = "\n".join(shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text and shape.text.strip())
            if text:
                sections.append({"text": text, "source_label": f"Slide {number}", "page_number": None, "slide_number": number})
    else:
        raise ValueError("Unsupported file. Upload a valid PDF or PPTX document.")
    if not sections:
        raise ValueError("No extractable text was found in this document.")
    return sections, kind


def _split(text: str) -> list[str]:
    words = _clean(text).split()
    step = max(1, CHUNK_SIZE - CHUNK_OVERLAP)
    return [" ".join(words[start : start + CHUNK_SIZE]) for start in range(0, len(words), step)]


def _chunks(sections: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for section in sections:
        for content in _split(section["text"]):
            result.append({"id": len(result), "content": content, "source_label": section["source_label"], "page_number": section["page_number"], "slide_number": section["slide_number"]})
    if not result:
        raise ValueError("No searchable chunks could be created from this document.")
    return result


def _model(device: str) -> SentenceTransformer:
    with _embedding_lock:
        if device not in _embedding_models:
            _embedding_models[device] = SentenceTransformer(EMBEDDING_MODEL, device=device)
    return _embedding_models[device]


def _normalize(vectors: Any) -> np.ndarray:
    array = np.asarray(vectors, dtype="float32")
    norms = np.linalg.norm(array, axis=1, keepdims=True)
    norms[norms == 0] = 1
    return (array / norms).astype("float32")


@spaces.GPU(duration=60)
def _encode_gpu(texts: list[str]) -> np.ndarray:
    return _normalize(_model("cuda").encode(texts, batch_size=8, show_progress_bar=False))


def _encode(texts: list[str]) -> np.ndarray:
    # CPU is available even when the shared ZeroGPU worker cannot start.
    # Serialize encoding and use small batches to bound peak model memory.
    with _encoding_lock:
        # Keep optional acceleration registered for ZeroGPU-hosted Spaces.
        # CPU is the default; a GPU worker failure still falls back safely.
        if os.environ.get("EMBEDDING_DEVICE", "cpu") == "cuda":
            try:
                return _encode_gpu(texts)
            except Exception:
                pass
        return _normalize(_model("cpu").encode(texts, batch_size=2, show_progress_bar=False))


def _prune(now: float | None = None) -> None:
    current = now or time.time()
    SESSION_ROOT.mkdir(parents=True, exist_ok=True)
    with _documents_lock:
        for key in [key for key, value in _documents.items() if current - value["last_accessed"] > DOCUMENT_TTL_SECONDS]:
            _documents.pop(key, None)
        while len(_documents) >= MAX_DOCUMENTS:
            _documents.popitem(last=False)
    session_dirs = sorted((path for path in SESSION_ROOT.iterdir() if path.is_dir()), key=lambda path: path.stat().st_mtime)
    for path in session_dirs:
        if current - path.stat().st_mtime > DOCUMENT_TTL_SECONDS:
            shutil.rmtree(path, ignore_errors=True)
    for path in session_dirs[:-MAX_DOCUMENTS]:
        shutil.rmtree(path, ignore_errors=True)


def _session_dir(document_id: str) -> Path:
    identifier = _clean(document_id)
    if not re.fullmatch(r"[a-f0-9]{32}", identifier):
        raise ValueError("Invalid temporary document ID.")
    return SESSION_ROOT / identifier


def _save_document(identifier: str, item: dict[str, Any]) -> None:
    directory = _session_dir(identifier)
    directory.mkdir(parents=True, exist_ok=True)
    metadata = {key: value for key, value in item.items() if key != "index"}
    (directory / "document.json").write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
    faiss.write_index(item["index"], str(directory / "index.faiss"))


def _load_document(identifier: str) -> dict[str, Any] | None:
    directory = _session_dir(identifier)
    metadata_path, index_path = directory / "document.json", directory / "index.faiss"
    if not metadata_path.is_file() or not index_path.is_file():
        return None
    item = json.loads(metadata_path.read_text(encoding="utf-8"))
    item["index"] = faiss.read_index(str(index_path))
    return item


def _document(document_id: str) -> dict[str, Any]:
    identifier = _clean(document_id)
    _prune()
    with _documents_lock:
        item = _documents.get(identifier) or _load_document(identifier)
        if not item:
            raise ValueError("This temporary document session expired. Prepare the document again.")
        item["last_accessed"] = time.time()
        _documents[identifier] = item
        _documents.move_to_end(identifier)
        os.utime(_session_dir(identifier), None)
        return item


def _error(exc: Exception) -> dict[str, Any]:
    return {"ok": False, "error": str(exc), "error_type": type(exc).__name__}


def prepare_document(file_path: str, original_filename: str = "") -> dict[str, Any]:
    try:
        sections, kind = _extract_sections(file_path)
        chunks = _chunks(sections)
        embeddings = _encode([item["content"] for item in chunks])
        index = faiss.IndexFlatIP(embeddings.shape[1])
        index.add(embeddings)
        identifier, now = uuid.uuid4().hex, time.time()
        filename = Path(_clean(original_filename) or file_path).name
        record = {"filename": filename, "document_type": kind, "section_count": len(sections), "chunk_count": len(chunks), "chunks": chunks, "index": index, "last_accessed": now, "quiz": None}
        _prune(now)
        with _documents_lock:
            _documents[identifier] = record
        _save_document(identifier, record)
        return {"ok": True, "document_id": identifier, "filename": record["filename"], "document_type": kind, "section_count": len(sections), "chunk_count": len(chunks), "expires_in_seconds": DOCUMENT_TTL_SECONDS}
    except Exception as exc:
        return _error(exc)


def _retrieve(document_id: str, question: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    item, query = _document(document_id), _clean(question)
    if not query:
        raise ValueError("Enter a question first.")
    embedding = _encode([query])
    scores, positions = item["index"].search(embedding, min(TOP_K, len(item["chunks"])))
    sources = []
    for position, score in zip(positions[0], scores[0]):
        if position >= 0:
            chunk = item["chunks"][int(position)]
            sources.append({"chunk_id": chunk["id"], "source_label": chunk["source_label"], "page_number": chunk["page_number"], "slide_number": chunk["slide_number"], "score": round(float(score), 4), "excerpt": chunk["content"][:900]})
    return item, sources


def _client() -> InferenceClient:
    token = os.environ.get("HF_TOKEN", "").strip()
    if not token:
        raise RuntimeError("The Space generation service is not configured.")
    return InferenceClient(provider="auto", api_key=token)


def _complete(system: str, prompt: str, max_tokens: int, temperature: float = 0.15) -> str:
    result = _client().chat.completions.create(model=GENERATION_MODEL, messages=[{"role": "system", "content": system}, {"role": "user", "content": prompt}], temperature=temperature, max_tokens=max_tokens)
    answer = result.choices[0].message.content
    if not answer:
        raise RuntimeError("The generation model returned an empty response.")
    return answer.strip()


def _evidence(sources: list[dict[str, Any]]) -> str:
    return "\n\n".join(f"[Source {number}: {source['source_label']}]\n{source['excerpt']}" for number, source in enumerate(sources, 1))


def ask_document(document_id: str, question: str) -> dict[str, Any]:
    try:
        _, sources = _retrieve(document_id, question)
        answer = _complete("Answer only from the document evidence. Never use outside knowledge. If evidence is insufficient, say it was not found. Be concise and cite page or slide labels.", f"Evidence:\n{_evidence(sources)}\n\nQuestion: {_clean(question)}", 500)
        return {"ok": True, "answer": answer, "sources": sources}
    except Exception as exc:
        return _error(exc)


def _context(item: dict[str, Any]) -> str:
    return "\n\n".join(f"[{chunk['source_label']}]\n{chunk['content']}" for chunk in item["chunks"])[:MAX_GENERATION_CONTEXT_CHARS]


def summarize_document(document_id: str) -> dict[str, Any]:
    try:
        summary = _complete("Create a clear study summary using only the document. Use short headings and bullets, preserve page or slide provenance, and add no facts.", f"Document:\n{_context(_document(document_id))}", 900)
        return {"ok": True, "summary": summary}
    except Exception as exc:
        return _error(exc)


def _json_array(text: str) -> list[dict[str, Any]]:
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip(), flags=re.IGNORECASE)
    try:
        value = json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\[[\s\S]*\]", cleaned)
        if not match:
            raise ValueError("The model did not return valid structured study material.")
        value = json.loads(match.group(0))
    if not isinstance(value, list):
        raise ValueError("The generated study material has an invalid format.")
    return [entry for entry in value if isinstance(entry, dict)]


def generate_flashcards(document_id: str, count: float = 6) -> dict[str, Any]:
    try:
        amount, item = max(3, min(10, int(count))), _document(document_id)
        raw = _complete("Create flashcards using only the document. Return a valid JSON array with no markdown.", f"Document:\n{_context(item)}\n\nCreate {amount} cards as [{'{'}\"question\":\"...\",\"answer\":\"...\",\"source\":\"Page/Slide ...\"{'}'}].", 1100, 0.1)
        cards = [{"question": _clean(x.get("question")), "answer": _clean(x.get("answer")), "source": _clean(x.get("source"))} for x in _json_array(raw) if _clean(x.get("question")) and _clean(x.get("answer"))]
        if not cards:
            raise ValueError("No valid flashcards were generated.")
        return {"ok": True, "flashcards": cards}
    except Exception as exc:
        return _error(exc)


def _normalize_quiz(entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    quiz = []
    for entry in entries:
        question = _clean(entry.get("question"))
        options = [_clean(value) for value in entry.get("options", []) if _clean(value)]
        answer = _clean(entry.get("answer"))
        if answer not in options:
            matches = [option for option in options if option[:1].upper() == answer[:1].upper()]
            answer = matches[0] if matches else answer
        if question and len(options) >= 2 and answer in options:
            quiz.append({"question": question, "options": options, "answer": answer, "source": _clean(entry.get("source"))})
    return quiz


def generate_quiz(document_id: str, count: float = 5) -> dict[str, Any]:
    try:
        amount, item = max(3, min(8, int(count))), _document(document_id)
        raw = _complete("Create a multiple-choice quiz using only the document. Return valid JSON only. The answer must exactly equal one complete option.", f"Document:\n{_context(item)}\n\nCreate {amount} questions as [{'{'}\"question\":\"...\",\"options\":[\"A. ...\",\"B. ...\",\"C. ...\",\"D. ...\"],\"answer\":\"A. ...\",\"source\":\"Page/Slide ...\"{'}'}].", 1500, 0.1)
        quiz = _normalize_quiz(_json_array(raw))
        if not quiz:
            raise ValueError("No valid quiz questions were generated.")
        item["quiz"] = quiz
        _save_document(_clean(document_id), item)
        return {"ok": True, "quiz": [{"id": i, "question": x["question"], "options": x["options"], "source": x["source"]} for i, x in enumerate(quiz)]}
    except Exception as exc:
        return _error(exc)


def grade_quiz(document_id: str, answers_json: str) -> dict[str, Any]:
    try:
        item = _document(document_id)
        if not item.get("quiz"):
            raise ValueError("Generate a quiz before submitting answers.")
        answers = json.loads(answers_json or "[]")
        if not isinstance(answers, list):
            raise ValueError("Quiz answers must be a JSON array.")
        results, correct = [], 0
        for index, entry in enumerate(item["quiz"]):
            selected = _clean(answers[index] if index < len(answers) else "")
            is_correct = bool(selected) and selected.casefold() == entry["answer"].casefold()
            correct += int(is_correct)
            results.append({"question": entry["question"], "selected_answer": selected, "correct_answer": entry["answer"], "is_correct": is_correct, "source": entry["source"]})
        total = len(results)
        return {"ok": True, "correct": correct, "total": total, "percent": round(correct * 100 / total) if total else 0, "results": results}
    except Exception as exc:
        return _error(exc)


def delete_document(document_id: str) -> dict[str, Any]:
    identifier = _clean(document_id)
    with _documents_lock:
        removed = _documents.pop(identifier, None)
    directory = _session_dir(identifier)
    existed = directory.exists()
    shutil.rmtree(directory, ignore_errors=True)
    return {"ok": True, "deleted": bool(removed) or existed}


with gr.Blocks(title="Smart Study Assistant") as demo:
    gr.Markdown("""
# Smart Study Assistant — cloud demo
Prepare a text-based PDF or PPTX once, then reuse its BGE-M3 embeddings for grounded chat and study tools.

**Pipeline:** PDF/PPTX → 500/75 chunks → normalized BGE-M3 → FAISS IndexFlatIP → hosted Qwen3 → answer + page/slide evidence

> Files are temporary, expire after 45 minutes of inactivity, and may disappear when the Space restarts. Do not upload confidential, sensitive, or personally identifiable information. Limits: 15 MB, 60 PDF pages, or 60 PowerPoint slides.
""")
    # API clients can upload Blobs without an extension; validate file signatures
    # in _extract_sections instead of rejecting valid PDF/PPTX uploads here.
    upload = gr.File(label="PDF or PPTX", file_types=["file"], type="filepath")
    original_filename = gr.Textbox(label="Original filename")
    document_id = gr.Textbox(label="Temporary document ID")
    question = gr.Textbox(label="Question")
    answers = gr.Textbox(label="Quiz answers JSON")
    output = gr.JSON(label="API result")
    with gr.Row():
        gr.Button("Prepare", variant="primary").click(prepare_document, [upload, original_filename], output, api_name="prepare_document")
        gr.Button("Ask").click(ask_document, [document_id, question], output, api_name="ask_document")
        gr.Button("Summary").click(summarize_document, document_id, output, api_name="summarize_document")
    with gr.Row():
        gr.Button("Flashcards").click(generate_flashcards, [document_id, gr.Number(value=6, visible=False)], output, api_name="generate_flashcards")
        gr.Button("Quiz").click(generate_quiz, [document_id, gr.Number(value=5, visible=False)], output, api_name="generate_quiz")
        gr.Button("Grade").click(grade_quiz, [document_id, answers], output, api_name="grade_quiz")
        gr.Button("Delete").click(delete_document, document_id, output, api_name="delete_document")


if __name__ == "__main__":
    demo.queue(default_concurrency_limit=4).launch()
