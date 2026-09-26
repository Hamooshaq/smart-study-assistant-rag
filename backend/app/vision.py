import base64
from dataclasses import dataclass
from pathlib import Path
import shutil
import subprocess
import tempfile

import requests

from .config import settings
from .extraction import ExtractedSection


@dataclass(frozen=True)
class ExtractedImage:
    path: Path
    source_type: str
    page_number: int | None = None
    slide_number: int | None = None


class VisionError(RuntimeError):
    pass


def extract_visual_sections(
    path: Path,
    extension: str,
    user_id: int,
    document_id: int,
    enabled: bool | None = None,
) -> list[ExtractedSection]:
    if enabled is None:
        enabled = settings.vlm_enabled
    if not enabled:
        return []
    images = extract_document_images(path, extension, user_id, document_id)
    sections: list[ExtractedSection] = []
    for image in images[: settings.vlm_max_images_per_document]:
        caption = describe_image(image.path)
        source_number = image.page_number if image.source_type == "pdf" else image.slide_number
        source_name = "Halaman" if image.source_type == "pdf" else "Slide"
        sections.append(
            ExtractedSection(
                text=f"Deskripsi visual dari {source_name} {source_number}: {caption}",
                source_type=image.source_type,
                page_number=image.page_number,
                slide_number=image.slide_number,
                is_visual=True,
            )
        )
    return sections


def extract_document_images(path: Path, extension: str, user_id: int, document_id: int) -> list[ExtractedImage]:
    if extension == ".pdf":
        return extract_pdf_images(path, user_id, document_id)
    if extension == ".pptx":
        return extract_pptx_images(path, user_id, document_id)
    if extension == ".ppt":
        with tempfile.TemporaryDirectory() as tmp:
            return extract_pptx_images(_convert_ppt_to_pptx(path, Path(tmp)), user_id, document_id)
    return []


def _convert_ppt_to_pptx(path: Path, output_dir: Path) -> Path:
    soffice = shutil.which("soffice")
    if not soffice:
        raise VisionError("File .ppt membutuhkan LibreOffice/soffice untuk ekstraksi gambar.")
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pptx", "--outdir", str(output_dir), str(path)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    converted = output_dir / f"{path.stem}.pptx"
    if converted.exists():
        return converted
    candidates = list(output_dir.glob("*.pptx"))
    if not candidates:
        raise VisionError("Konversi .ppt ke .pptx gagal.")
    return candidates[0]


def _image_output_dir(user_id: int, document_id: int) -> Path:
    out = settings.vision_dir / f"user_{user_id}" / f"document_{document_id}"
    out.mkdir(parents=True, exist_ok=True)
    return out


def extract_pdf_images(path: Path, user_id: int, document_id: int) -> list[ExtractedImage]:
    try:
        import fitz
    except ImportError as exc:
        raise VisionError("PyMuPDF belum terpasang. Jalankan pip install -r requirements.txt.") from exc

    out_dir = _image_output_dir(user_id, document_id)
    images: list[ExtractedImage] = []
    with fitz.open(path) as doc:
        for page_index in range(len(doc)):
            page = doc[page_index]
            for image_index, image_info in enumerate(page.get_images(full=True), start=1):
                xref = image_info[0]
                data = doc.extract_image(xref)
                target = out_dir / f"page_{page_index + 1}_image_{image_index}.{data.get('ext', 'png')}"
                target.write_bytes(data["image"])
                images.append(ExtractedImage(path=target, source_type="pdf", page_number=page_index + 1))
    return images


def extract_pptx_images(path: Path, user_id: int, document_id: int) -> list[ExtractedImage]:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise VisionError("python-pptx belum terpasang. Jalankan pip install -r requirements.txt.") from exc

    out_dir = _image_output_dir(user_id, document_id)
    images: list[ExtractedImage] = []
    presentation = Presentation(str(path))
    for slide_number, slide in enumerate(presentation.slides, start=1):
        picture_index = 0
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            picture_index += 1
            target = out_dir / f"slide_{slide_number}_image_{picture_index}.{shape.image.ext or 'png'}"
            target.write_bytes(shape.image.blob)
            images.append(ExtractedImage(path=target, source_type="ppt", slide_number=slide_number))
    return images


def describe_image(path: Path) -> str:
    payload = {
        "model": settings.vlm_model,
        "messages": [
            {
                "role": "user",
                "content": (
                    "Jelaskan isi gambar ini untuk kebutuhan belajar mahasiswa. "
                    "Jika ada diagram, flowchart, tabel, screenshot, ilustrasi, atau teks penting, tuliskan poin-poinnya. "
                    "Jawab dalam bahasa Indonesia dan jangan menambahkan informasi di luar gambar."
                ),
                "images": [base64.b64encode(path.read_bytes()).decode("utf-8")],
            }
        ],
        "stream": False,
        "options": {"temperature": 0.1},
    }
    try:
        response = requests.post(f"{settings.ollama_url.rstrip('/')}/api/chat", json=payload, timeout=180)
    except requests.RequestException as exc:
        raise VisionError("Ollama VLM belum bisa dihubungi. Matikan analisis gambar atau jalankan model vision di Ollama.") from exc
    if response.status_code >= 400:
        raise VisionError(f"Ollama VLM error {response.status_code}: {response.text[:300]}")
    return response.json().get("message", {}).get("content", "").strip()
