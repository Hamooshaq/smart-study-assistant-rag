from dataclasses import dataclass
from pathlib import Path
import shutil
import subprocess
import tempfile


@dataclass(frozen=True)
class ExtractedSection:
    text: str
    source_type: str
    page_number: int | None = None
    slide_number: int | None = None
    is_visual: bool = False


def extract_document_text(path: Path, extension: str) -> list[ExtractedSection]:
    if extension == ".pdf":
        return extract_pdf_text(path)
    if extension in {".ppt", ".pptx"}:
        return extract_presentation_text(path, extension)
    raise ValueError("Format file tidak didukung.")


def extract_pdf_text(path: Path) -> list[ExtractedSection]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("pdfplumber belum terpasang. Jalankan pip install -r requirements.txt.") from exc

    sections: list[ExtractedSection] = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(ExtractedSection(text=text, source_type="pdf", page_number=page_number))
    if not sections:
        raise ValueError("PDF tidak memiliki teks yang bisa diekstrak. OCR belum termasuk pipeline utama.")
    return sections


def extract_presentation_text(path: Path, extension: str) -> list[ExtractedSection]:
    cleanup_dir: tempfile.TemporaryDirectory[str] | None = None
    pptx_path = path
    if extension == ".ppt":
        cleanup_dir = tempfile.TemporaryDirectory()
        pptx_path = _convert_ppt_to_pptx(path, Path(cleanup_dir.name))
    try:
        return _extract_pptx_text(pptx_path)
    finally:
        if cleanup_dir:
            cleanup_dir.cleanup()


def _convert_ppt_to_pptx(path: Path, output_dir: Path) -> Path:
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError("File .ppt membutuhkan LibreOffice/soffice untuk konversi ke .pptx.")
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pptx", "--outdir", str(output_dir), str(path)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    converted = output_dir / f"{path.stem}.pptx"
    if converted.exists():
        return converted
    candidates = list(output_dir.glob("*.pptx"))
    if not candidates:
        raise RuntimeError("Konversi .ppt ke .pptx gagal.")
    return candidates[0]


def _extract_pptx_text(path: Path) -> list[ExtractedSection]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx belum terpasang. Jalankan pip install -r requirements.txt.") from exc

    presentation = Presentation(str(path))
    sections: list[ExtractedSection] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text.strip())
        text = "\n".join(t for t in texts if t).strip()
        if text:
            sections.append(ExtractedSection(text=text, source_type="ppt", slide_number=slide_number))
    if not sections:
        raise ValueError("PowerPoint tidak memiliki teks yang bisa diekstrak.")
    return sections
